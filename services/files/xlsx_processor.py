"""XLSX hajmi (qatorlar soni)."""

from openpyxl import load_workbook

from data import config
from services.files.base import BaseProcessor, FileAnalysis, ceil_div


class XlsxProcessor(BaseProcessor):
    unit_name = "qator"

    def analyze(self, path: str) -> FileAnalysis:
        wb = load_workbook(path, read_only=True)
        rows = sum(ws.max_row or 0 for ws in wb.worksheets)
        wb.close()
        rows = max(1, rows)
        pages = max(1, ceil_div(rows, config.ROWS_PER_PAGE))
        return FileAnalysis(unit_count=rows, unit_name=self.unit_name, pages=pages)
